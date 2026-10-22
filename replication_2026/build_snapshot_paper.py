from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent
doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
st=doc.styles['Normal'];st.font.size=Pt(11);st.paragraph_format.line_spacing=1.03;st.paragraph_format.space_after=Pt(6)
for name,size in [('Title',20),('Heading 1',12),('Heading 2',11)]:
    st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(5)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    if style.name in ['Normal','Title','Heading 1','Heading 2']:
        fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            if qn('w:'+key) in fonts.attrib:del fonts.attrib[qn('w:'+key)]
        for key in ['ascii','hAnsi']:fonts.set(qn('w:'+key),'Times New Roman')
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
lines=(ROOT/'SNAPSHOT_OFI_PAPER.md').read_text().splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line=='<!-- pagebreak -->':doc.add_page_break();i+=1;continue
    if line.startswith('|'):
        group=[]
        while i<len(lines) and lines[i].startswith('|'):
            if not re.fullmatch(r'[| :\-]+',lines[i]):group.append([v.strip() for v in lines[i].strip('|').split('|')])
            i+=1
        t=doc.add_table(rows=0,cols=len(group[0]));t.autofit=False;t.columns[0].width=Inches(5.8);t.columns[1].width=Inches(1.2)
        for j,row in enumerate(group):
            cells=t.add_row().cells
            for k,val in enumerate(row):
                cells[k].text=val
                for p in cells[k].paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                    for r in p.runs:r.font.size=Pt(10);r.bold=j==0
            if j==0:
                pr=t.rows[j]._tr.get_or_add_trPr();tag=OxmlElement('w:tblHeader');pr.append(tag)
        continue
    if line.startswith('!['):
        path=re.search(r'\]\((.*?)\)',line).group(1);p=doc.add_paragraph();p.add_run().add_picture(str(ROOT/path),width=Inches(6.8));p.paragraph_format.space_after=Pt(3)
    elif line.startswith('$$'):
        p=doc.add_paragraph(line.strip('$'));p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(4)
        for r in p.runs:r.font.name='Cambria Math';r.font.size=Pt(10)
    elif line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
    elif line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 1')
    elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2')
    else:
        p=doc.add_paragraph();pieces=re.split(r'(\*[^*]+\*)',line)
        for piece in pieces:
            r=p.add_run(piece.strip('*') if piece.startswith('*') else piece);r.italic=piece.startswith('*')
        if line.startswith('*Figure'):
            for r in p.runs:r.font.size=Pt(9)
        if line.startswith('[') or line.startswith('Code, full'):
            for r in p.runs:r.font.size=Pt(9)
    i+=1
doc.core_properties.title='What snapshot sampling does to order flow imbalance';doc.core_properties.author='Brendan Connelly and Rigel Mummers';doc.save(ROOT/'SNAPSHOT_OFI_PAPER.docx')
print('Created editable paper')
